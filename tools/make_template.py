import json, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
d=json.load(open('/tmp/claude-0/tpl.json'))
NAVY='16324F'
wb=openpyxl.Workbook(); ws=wb.active; ws.title='Instructions'
hdr=Font(bold=True,color='FFFFFF'); fill=PatternFill('solid',fgColor=NAVY); thin=Side(style='thin',color='C9CED6')
def title(ws,t,sub):
    ws['A1']=t; ws['A1'].font=Font(bold=True,size=16,color=NAVY); ws['A2']=sub; ws['A2'].font=Font(italic=True,color='5D6670')
title(ws,'Family Ready — import template','Fill in the other sheets, then in the app go to Settings (⚙️) → Excel → Import.')
lines=['How to use',
'1. Fill in the Family, Contacts, Suppliers, Supplies and Places sheets — one row per person, contact, supplier, item or place.',
'2. Keep the heading row (row 4) exactly as it is. Use the drop-down lists where there is one.',
'3. Grey rows starting with "e.g." are examples — overwrite or delete them (they are ignored on import).',
'4. Leave a sheet empty if you don\'t need it.',
'5. In the app: Settings (⚙️) → Excel — import & export → Import Excel / CSV → choose this file.',
'6. Rows with the same Name as something already in the app are updated; new names are added. Nothing is deleted unless you tick "Replace".',
'7. Supplies: put Y in "Have it" for things you have, and a use-by date. The app reminds you before things expire.',
'8. You can also Export from the app, edit that file in Excel, and import it again.',
'','Tips','• Phone numbers: type them as text, e.g. 021 123 4567 (start with an apostrophe if Excel drops the leading 0).','• Dates: any normal date format, e.g. 29/10/2026.','• Places and suppliers: latitude/longitude are optional — you can pin locations on the map in the app instead.']
for i,l in enumerate(lines, start=4):
    ws.cell(row=i,column=1,value=l)
    if l in ('How to use','Tips'): ws.cell(row=i,column=1).font=Font(bold=True,size=12,color=NAVY)
ws.column_dimensions['A'].width=120
L=wb.create_sheet('Lists')
for i,a in enumerate(d['ages'],1): L.cell(row=i,column=1,value=a)
for i,c in enumerate(d['cats'],1): L.cell(row=i,column=2,value=c[0])
for i,v in enumerate(['Y','N'],1): L.cell(row=i,column=3,value=v)
places=['Water source','Shelter','Supply cache','Danger / avoid','Food / fruit trees','Other']
for i,v in enumerate(places,1): L.cell(row=i,column=4,value=v)
L.sheet_state='hidden'
def sheet(name, sub, cols, examples, validations=(), rows=()):
    ws=wb.create_sheet(name); title(ws,name,sub)
    for j,(h,w) in enumerate(cols,1):
        c=ws.cell(row=4,column=j,value=h); c.font=hdr; c.fill=fill; c.alignment=Alignment(wrap_text=True,vertical='center')
        ws.column_dimensions[get_column_letter(j)].width=w
    ws.row_dimensions[4].height=32; ws.freeze_panes='B5'
    r=5
    for ex in examples:
        for j,v in enumerate(ex,1): ws.cell(row=r,column=j,value=v).font=Font(italic=True,color='8A939D')
        r+=1
    for row in rows:
        for j,v in enumerate(row,1): ws.cell(row=r,column=j,value=v)
        r+=1
    for col,ref in validations:
        dv=DataValidation(type='list',formula1=ref,allow_blank=True); ws.add_data_validation(dv); dv.add(f'{col}5:{col}500')
    for rr in ws.iter_rows(min_row=5,max_row=max(r,40),max_col=len(cols)):
        for c in rr: c.alignment=Alignment(wrap_text=True,vertical='top'); c.border=Border(bottom=thin)
    ws.auto_filter.ref=f"A4:{get_column_letter(len(cols))}{max(r,40)}"
sheet('Family','Everyone in the household. Age group sets food and protein amounts.',
 [('Name',16),('Who they are',16),('Age group',20),('Mobile',15),('Date of birth',13),('Medical conditions',24),('Medicines, doses and times',28),('Allergies',18),('Blood type',9),('GP / doctor and phone',22),('Glasses, hearing aids, other aids',22),('Mobility needs',22),('Who helps them evacuate',20),('Notes',24)],
 [['e.g. Jane Smith','Mum','Woman 19–70','021 000 0000','','Asthma','Ventolin 2 puffs as needed','Penicillin','O+','Dr Lee, 09 000 0000','Glasses','','','']], [('C',f"=Lists!$A$1:$A${len(d['ages'])}")])
sheet('Contacts','Neighbours, relatives, school, work — and one person outside Auckland.',
 [('Name',20),('Who they are',20),('Phone',16),('Other phone / email',24),('Address',26),('Out of town (Y/N)',10),('Notes',30)],
 [['e.g. Aunty Mere','Aunt (Hamilton)','07 000 0000','mere@example.com','Hamilton','Y','Our out-of-town check-in person']], [('F','=Lists!$C$1:$C$2')])
sheet('Suppliers','People and places to get food, animals and essentials from when the shops are empty.',
 [('Name',22),('Category',20),('What they supply',30),('Phone',15),('Other contact',20),('Address / where',22),('Would take in trade',20),('Last checked',12),('Notes',26),('Latitude',11),('Longitude',11)],
 [["e.g. Bob's hens",'Eggs & chickens','Eggs, point-of-lay hens, layer pellets','09 000 0000','Facebook: Bob Hens','Oratia','Cash or firewood','','','','']], [('B',f"=Lists!$B$1:$B${len(d['cats'])}")])
sheet('Supplies',"Every item in the app's supplies list. Put Y if we have it, and a use-by date. Add your own items at the bottom.",
 [('Item',48),('Group',18),('How much we need',16),('Have it (Y/N)',10),('Use-by date',13),('Notes',30)], [], [('D','=Lists!$C$1:$C$2')],
 rows=[[s[1],s[0],'','','',''] for s in d['supplies']])
sheet('Places','Map places: water sources, shelters, caches, things to avoid.',
 [('Name',26),('Type',18),('Latitude',12),('Longitude',12),('Notes',34)],
 [["e.g. Neighbour's rain tank",'Water source',-36.91,174.65,'Ask first']], [('B',f"=Lists!$D$1:$D${len(places)}")])
wb.move_sheet('Lists', offset=10)
wb.save('/home/claude/ready/dist/family-ready-template.xlsx'); print('ok')
